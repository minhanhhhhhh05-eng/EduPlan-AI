
from docx import Document
from docx.shared import Pt, Cm
from docx.shared import RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from io import BytesIO


def tao_file_word(thong_tin, giao_an):
    doc = Document()

    # Thiet lap kho giay A4 va can le
    section = doc.sections[0]

    section.page_width = Cm(21)
    section.page_height = Cm(29.7)

    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)

    # Định dạng văn bản
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(13)

    # Dinh dang gian dong va khoang cach doan
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(6)

    # Dinh dang cac tieu de Word
    for ten_style in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        tieu_de = doc.styles[ten_style]
        tieu_de.font.name = "Times New Roman"
        tieu_de.element.rPr.rFonts.set(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia",
            "Times New Roman"
        )
        
        for thuoc_tinh in ["ascii", "hAnsi", "cs"]:
            tieu_de.element.rPr.rFonts.set(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}" + thuoc_tinh,
                "Times New Roman"
            )

        # Loai bo tham chieu den font chu de Word
        rfonts = tieu_de.element.rPr.rFonts

        for ten_font in ["asciiTheme", "hAnsiTheme",
                         "eastAsiaTheme", "csTheme"]:
            rfonts.attrib.pop(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}" + ten_font,
                None
            )

        tieu_de.font.bold = True
        tieu_de.font.color.rgb = RGBColor(0, 0, 0)

    doc.styles["Title"].font.size = Pt(16)
    doc.styles["Heading 1"].font.size = Pt(14)
    doc.styles["Heading 2"].font.size = Pt(13)
    doc.styles["Heading 3"].font.size = Pt(13)

    # Thu gon khoang cach giua cac tieu de
    for ten_style in ["Heading 1", "Heading 2", "Heading 3"]:
        dinh_dang = doc.styles[ten_style].paragraph_format
        dinh_dang.space_before = Pt(6)
        dinh_dang.space_after = Pt(3)
        dinh_dang.line_spacing = 1.15     
        dinh_dang.keep_with_next = True
        dinh_dang.keep_together = True
        dinh_dang.page_break_before = False

    # Thông tin bài dạy
    
    # Tieu de chinh cua ke hoach bai day
    tieu_de = doc.add_paragraph()
    tieu_de.alignment = WD_ALIGN_PARAGRAPH.CENTER

    chu = tieu_de.add_run("KẾ HOẠCH BÀI DẠY")
    chu.bold = True
    chu.font.name = "Times New Roman"
    chu.font.size = Pt(16)
    chu.font.color.rgb = RGBColor(0, 0, 0)


    for ten, gia_tri in thong_tin.items():
        doc.add_paragraph(f"{ten}: {gia_tri}")

    # Ba phần chính của giáo án
    
    for ten_phan, noi_dung in giao_an.items():
        doc.add_heading(ten_phan, level=1)

        if isinstance(noi_dung, dict):
            for ten_muc, chi_tiet in noi_dung.items():
                doc.add_heading(ten_muc, level=2)

                if isinstance(chi_tiet, dict):
                    for muc, gia_tri in chi_tiet.items():
                        co_noi_dung = (
                            any(str(v).strip() for v in gia_tri.values())
                            if isinstance(gia_tri, dict)
                            else bool(str(gia_tri).strip())
                        )

                        tieu_de_muc = doc.add_heading(muc, level=3)
                        tieu_de_muc.paragraph_format.keep_with_next = co_noi_dung

                        if isinstance(gia_tri, dict):
                            for buoc, noi_dung_buoc in gia_tri.items():
                                if str(noi_dung_buoc).strip():
                                    doc.add_paragraph(f"{buoc}:")
                                    for dong in str(noi_dung_buoc).splitlines():
                                        if dong.strip():
                                            doc.add_paragraph(dong.strip())
                        else:
                            for dong in str(gia_tri).splitlines():
                                if dong.strip():
                                    doc.add_paragraph(dong.strip())
                else:
                    if str(chi_tiet).strip():
                        doc.add_paragraph(str(chi_tiet))
        else:
            if str(noi_dung).strip():
                doc.add_paragraph(str(noi_dung))

    # Lưu Word vào bộ nhớ
    file_word = BytesIO()
    doc.save(file_word)
    file_word.seek(0)

    return file_word

# Kiểm thử chức năng xuất Word
if __name__ == "__main__":
    from giao_an import tao_khung_giao_an

    thong_tin = {
        "Trường": "Trường THCS và THPT Nguyễn Văn Hai",
        "Tổ chuyên môn": "Tổ Tiếng Anh",
        "Giáo viên": "Giáo viên kiểm thử",
        "Môn học": "Tiếng Anh",
        "Lớp": "11",
        "Tên bài": "Unit 3 - Getting Started",
        "Số tiết": "1"
    }

    giao_an = tao_khung_giao_an()

    giao_an["I. MỤC TIÊU"]["Kiến thức"] = (
        "Nhận biết nội dung chính của bài học."
    )

    giao_an["I. MỤC TIÊU"]["Năng lực"] = (
        "Phát triển năng lực giao tiếp và hợp tác."
    )

    giao_an["I. MỤC TIÊU"]["Phẩm chất"] = (
        "Hình thành tinh thần trách nhiệm trong học tập."
    )

    giao_an["II. THIẾT BỊ DẠY HỌC VÀ HỌC LIỆU"] = (
        "SGK, máy tính, máy chiếu và phiếu học tập."
    )

    file_word = tao_file_word(thong_tin, giao_an)

    with open("Kiem_thu_EduPlan_AI.docx", "wb") as f:
        f.write(file_word.getvalue())

    print("DA TAO FILE WORD THANH CONG!")
